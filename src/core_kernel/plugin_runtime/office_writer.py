"""Render an Office outline to real .docx / .pptx with a shared design system."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from src.core_kernel.plugin_runtime.office_math import (
    equation_fallback_lines,
    iter_text_math_parts,
    latex_from_block,
    latex_to_omml_element,
    latex_to_unicode,
    text_has_math,
)
from src.core_kernel.plugin_runtime.office_outline import DEFAULT_THEME, OfficeOutlineError
from src.core_kernel.plugin_runtime.office_store import copy_asset
from src.core_kernel.plugin_runtime.office_style import (
    layout_from_style,
    normalize_style_id,
    numbered_caption,
    theme_from_style,
    token_color,
)

_REPO_ROOT = Path(__file__).resolve().parents[3]
IMAGE_DIR = _REPO_ROOT / "data" / "generated_images"


def _pack_of(outline: Dict[str, Any]) -> Tuple[str, Dict[str, str], Dict[str, Any]]:
    sid = normalize_style_id(outline.get("style_id"))
    theme = {**DEFAULT_THEME, **theme_from_style(sid)}
    return sid, theme, layout_from_style(sid)


def _hex_rgb(value: str) -> Tuple[int, int, int]:
    text = str(value or "").strip().lstrip("#")
    if len(text) == 3:
        text = "".join(c * 2 for c in text)
    if len(text) != 6:
        text = "1C2430"
    return int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)


def resolve_image_file(
    item: Dict[str, Any],
    *,
    cwd: str = "",
    doc_id: str = "",
) -> Tuple[Optional[Path], str]:
    """Return (local_path, preview_url). Mutates item['preview_url'] when resolved."""
    url = str(item.get("url") or "").strip()
    path_s = str(item.get("path") or item.get("src") or "").strip()
    preview = str(item.get("preview_url") or "").strip()

    if url.startswith("/api/generated-images/"):
        name = Path(url.split("/", 3)[-1] if url.count("/") >= 3 else url).name
        local = IMAGE_DIR / name
        if local.is_file():
            item["preview_url"] = url
            return local, url
    if url.startswith("/api/office/assets/"):
        from src.core_kernel.plugin_runtime.office_store import resolve_asset_file

        name = Path(url).name
        local = resolve_asset_file(name)
        if local:
            item["preview_url"] = url
            return local, url

    candidate: Optional[Path] = None
    if path_s:
        raw = Path(path_s)
        if raw.is_file():
            candidate = raw
        elif cwd:
            try:
                from src.adapters.workspaces.store import resolve_under_workspace

                hit = resolve_under_workspace(cwd, path_s)
                if hit.is_file():
                    candidate = hit
            except Exception:
                fallback = Path(cwd) / path_s
                if fallback.is_file():
                    candidate = fallback

    if candidate and candidate.is_file():
        try:
            dest, preview_url = copy_asset(candidate, doc_id=doc_id or "off")
            item["preview_url"] = preview_url
            return dest, preview_url
        except Exception:
            item["preview_url"] = preview or url
            return candidate, preview or url

    if url.startswith("http://") or url.startswith("https://"):
        item["preview_url"] = url
        return None, url
    item["preview_url"] = preview or url
    return None, preview or url


def _try_omml(paragraph, latex: str, *, display: bool) -> bool:
    el = latex_to_omml_element(latex, display=display)
    if el is None:
        return False
    try:
        paragraph._element.append(el)
        return True
    except Exception:
        return False


def _add_math_fallback_run(paragraph, latex: str, theme: Dict[str, str], *, size_pt: float):
    pretty, source = equation_fallback_lines(latex)
    run = paragraph.add_run(pretty or source)
    _set_run_font(run, theme, size_pt=size_pt, italic=True, color=theme.get("ink") or "#1C2430")
    try:
        run.font.name = "Cambria Math"
    except Exception:
        pass
    return run


def _fill_rich_text(
    paragraph,
    text: str,
    theme: Dict[str, str],
    *,
    size_pt: float,
    bold: bool = False,
    italic: bool = False,
    color: str = "",
    heading: bool = False,
):
    if not text_has_math(text):
        run = paragraph.add_run(text)
        _set_run_font(run, theme, size_pt=size_pt, bold=bold, italic=italic, color=color, heading=heading)
        return paragraph
    for kind, payload in iter_text_math_parts(text):
        if kind == "text":
            if not payload:
                continue
            run = paragraph.add_run(payload)
            _set_run_font(run, theme, size_pt=size_pt, bold=bold, italic=italic, color=color, heading=heading)
            continue
        if not _try_omml(paragraph, payload, display=False):
            _add_math_fallback_run(paragraph, payload, theme, size_pt=size_pt + (1 if kind == "math_block" else 0))
    return paragraph


def _add_equation_paragraph(
    doc,
    block: Dict[str, Any],
    theme: Dict[str, str],
    *,
    align_center=True,
    pack: Optional[Dict[str, Any]] = None,
    number: int = 0,
):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    flags = pack or {}
    latex = latex_from_block(block)
    display = str(block.get("display") or "block") != "inline"
    size = float(flags.get("equation_size_pt") or 14) if display else 12
    p = doc.add_paragraph()
    if display and align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p_spacing(p, before=80 if display else 20, after=40 if display else 80, line=1.2 if display else 1.35)
    if display and flags.get("equation_frame"):
        _para_borders(p, color=theme.get("rule") or "#C9D4CE", sz="8")
    if not _try_omml(p, latex, display=display):
        _add_math_fallback_run(p, latex, theme, size_pt=size)
        src = latex
        if src and display:
            s = doc.add_paragraph()
            if align_center:
                s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = s.add_run(src)
            _set_run_font(run, theme, size_pt=9, italic=True, color=theme["muted"])
            _p_spacing(s, before=0, after=60, line=1.15)
    if number and display:
        npara = doc.add_paragraph()
        npara.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        nr = npara.add_run(f"({number})")
        _set_run_font(nr, theme, size_pt=10, color=theme["muted"])
        _p_spacing(npara, before=0, after=20, line=1.1)
    caption = str(block.get("caption") or "").strip()
    if caption:
        c = doc.add_paragraph()
        if display and align_center:
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = c.add_run(caption)
        _set_run_font(cr, theme, size_pt=9.5, italic=True, color=theme["muted"])
        _p_spacing(c, before=0, after=80, line=1.15)
    return p


def _set_run_font(
    run,
    theme: Dict[str, str],
    *,
    size_pt: float,
    bold: bool = False,
    color: str = "",
    italic: bool = False,
    heading: bool = False,
):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    latin = (theme.get("font_heading") if heading else theme.get("font_body")) or "Calibri"
    run.font.name = latin
    r = _hex_rgb(color or theme.get("ink") or "#1C2430")
    run.font.color.rgb = RGBColor(*r)
    try:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        east = (
            theme.get("font_heading_east_asia")
            if heading
            else theme.get("font_east_asia")
        ) or theme.get("font_east_asia") or "微软雅黑"
        rFonts.set(qn("w:ascii"), latin)
        rFonts.set(qn("w:hAnsi"), latin)
        rFonts.set(qn("w:eastAsia"), east)
    except Exception:
        pass


def _p_spacing(paragraph, *, before: int = 0, after: int = 80, line: float = 1.28):
    pf = paragraph.paragraph_format
    from docx.shared import Pt

    pf.space_before = Pt(before / 20) if before else Pt(0)
    pf.space_after = Pt(after / 20) if after else Pt(8)
    try:
        pf.line_spacing = line
    except Exception:
        pass


def _para_borders(paragraph, *, color: str, sz: str = "8", sides: Tuple[str, ...] = ("top", "left", "bottom", "right"), space: str = "6"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), space)
        el.set(qn("w:color"), color.lstrip("#"))
        pBdr.append(el)
    pPr.append(pBdr)


def _bottom_border(paragraph, color: str, *, sz: str = "12"):
    _para_borders(paragraph, color=color, sz=sz, sides=("bottom",), space="6")


def _left_border(paragraph, color: str, *, sz: str = "24"):
    _para_borders(paragraph, color=color, sz=sz, sides=("left",), space="10")


def _shade_paragraph(paragraph, hex_color: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def _set_table_borders(table, color: str, *, sz: str = "4"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color.lstrip("#"))
        borders.append(el)
    tblPr.append(borders)


def _add_page_field(paragraph, theme: Dict[str, str]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    _set_run_font(run, theme, size_pt=9, color=theme.get("muted") or "#5C6B7A")
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(end)


def _shade_cell(cell, hex_color: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def render_docx(outline: Dict[str, Any], *, cwd: str = "") -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise OfficeOutlineError("python-docx is not installed") from exc

    _sid, theme, pack = _pack_of(outline)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    margin = float(pack.get("page_margin_in") or 0.95)
    section.left_margin = Inches(margin)
    section.right_margin = Inches(margin)
    section.top_margin = Inches(float(pack.get("page_margin_top_in") or 0.9))
    section.bottom_margin = Inches(float(pack.get("page_margin_bottom_in") or 0.85))

    body_pt = float(pack.get("body_size_pt") or 11)
    line = float(pack.get("line_spacing") or 1.35)

    try:
        normal = doc.styles["Normal"]
        normal.font.name = theme.get("font_body") or "Calibri"
        normal.font.size = Pt(body_pt)
        normal.font.color.rgb = RGBColor(*_hex_rgb(theme["ink"]))
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), theme.get("font_east_asia") or "微软雅黑")
    except Exception:
        pass

    if pack.get("word_header_title"):
        try:
            hp = section.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr = hp.add_run(str(outline.get("title") or ""))
            _set_run_font(hr, theme, size_pt=9, italic=True, color=theme["muted"])
            _bottom_border(hp, theme["rule"], sz="6")
        except Exception:
            pass
    if pack.get("word_footer_page"):
        try:
            fp = section.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_page_field(fp, theme)
        except Exception:
            pass

    doc_id = str(outline.get("doc_id") or "")
    fig_n = 0
    tbl_n = 0
    eq_n = 0
    for block in outline.get("blocks") or []:
        if not isinstance(block, dict):
            continue
        btype = block.get("type")
        if btype == "heading":
            level = int(block.get("level") or 1)
            p = doc.add_paragraph()
            if level == 1:
                _fill_rich_text(
                    p,
                    str(block.get("text") or ""),
                    theme,
                    size_pt=float(pack.get("h1_size_pt") or 26),
                    bold=True,
                    color=theme["ink"],
                    heading=True,
                )
                _p_spacing(p, before=160, after=80, line=1.15)
                underline = pack.get("h1_underline") or "thick"
                if underline == "thick":
                    _bottom_border(p, theme["accent"], sz="16")
                elif underline == "thin":
                    _bottom_border(p, theme["rule"], sz="6")
            elif level == 2:
                _fill_rich_text(
                    p,
                    str(block.get("text") or ""),
                    theme,
                    size_pt=float(pack.get("h2_size_pt") or 16),
                    bold=True,
                    color=token_color(theme, str(pack.get("h2_color") or "accent_dark")),
                    heading=True,
                )
                _p_spacing(p, before=140, after=60, line=1.2)
            else:
                _fill_rich_text(
                    p,
                    str(block.get("text") or ""),
                    theme,
                    size_pt=float(pack.get("h3_size_pt") or 13),
                    bold=True,
                    color=token_color(theme, str(pack.get("h3_color") or "ink_soft")),
                    heading=True,
                )
                _p_spacing(p, before=100, after=40, line=1.2)
        elif btype == "paragraph":
            p = doc.add_paragraph()
            _fill_rich_text(p, str(block.get("text") or ""), theme, size_pt=body_pt, color=theme["ink"])
            _p_spacing(p, before=20, after=80, line=line)
        elif btype == "equation":
            display = str(block.get("display") or "block") != "inline"
            number = 0
            if display and pack.get("equation_numbers"):
                eq_n += 1
                number = eq_n
            _add_equation_paragraph(doc, block, theme, pack=pack, number=number)
        elif btype in {"bullet_list", "numbered_list"}:
            style = "List Number" if btype == "numbered_list" else "List Bullet"
            for item in block.get("items") or []:
                p = doc.add_paragraph(style=style)
                _fill_rich_text(p, str(item), theme, size_pt=body_pt, color=theme["ink"])
                _p_spacing(p, before=20, after=40, line=min(line, 1.35))
        elif btype == "quote":
            p = doc.add_paragraph()
            _fill_rich_text(p, str(block.get("text") or ""), theme, size_pt=body_pt + 1, italic=True, color=theme["ink_soft"])
            _p_spacing(p, before=80, after=40, line=max(line, 1.4))
            indent = 0.4 if pack.get("quote_bar") == "hairline" else 0.28
            p.paragraph_format.left_indent = Inches(indent)
            bar = pack.get("quote_bar") or "accent"
            if bar == "accent":
                _left_border(p, theme["accent"], sz=str(pack.get("quote_bar_sz") or "24"))
            elif bar == "hairline":
                _left_border(p, theme["rule"], sz=str(pack.get("quote_bar_sz") or "6"))
            if pack.get("quote_shade"):
                _shade_paragraph(p, theme.get("quote_bg") or theme["paper"])
            attr = str(block.get("attribution") or "").strip()
            if attr:
                a = doc.add_paragraph()
                ar = a.add_run(f"— {attr}")
                _set_run_font(ar, theme, size_pt=10, italic=True, color=theme["muted"])
                a.paragraph_format.left_indent = Inches(indent)
                _p_spacing(a, before=0, after=80, line=1.2)
        elif btype == "table":
            headers = [str(h) for h in (block.get("headers") or [])]
            rows = block.get("rows") or []
            cols = max(len(headers), max((len(r) for r in rows), default=1), 1)
            table = doc.add_table(rows=1 + len(rows), cols=cols)
            table.style = "Table Grid"
            table.autofit = True
            header_fill = bool(pack.get("table_header_fill"))
            zebra = bool(pack.get("table_zebra"))
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = ""
                p = cell.paragraphs[0]
                fg = theme["header_fg"] if header_fill else theme["ink"]
                _fill_rich_text(p, h, theme, size_pt=10, bold=True, color=fg)
                _shade_cell(cell, theme["accent_dark"] if header_fill else theme["paper"])
            for ri, row in enumerate(rows, 1):
                if zebra:
                    bg = theme["paper"] if ri % 2 else theme["paper_alt"]
                else:
                    bg = theme["paper"]
                for ci in range(cols):
                    cell = table.rows[ri].cells[ci]
                    cell.text = ""
                    val = row[ci] if ci < len(row) else ""
                    p = cell.paragraphs[0]
                    _fill_rich_text(p, str(val), theme, size_pt=10, color=theme["ink"])
                    _shade_cell(cell, bg)
            if not header_fill:
                _set_table_borders(table, theme["rule"], sz="4")
            tbl_n += 1
            cap = numbered_caption("表", tbl_n, str(block.get("caption") or ""), enabled=bool(pack.get("caption_numbers")))
            if cap:
                c = doc.add_paragraph()
                cr = c.add_run(cap)
                _set_run_font(cr, theme, size_pt=9.5, italic=True, color=theme["muted"])
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _p_spacing(c, before=20, after=80, line=1.2)
            else:
                doc.add_paragraph()
        elif btype == "page_break":
            doc.add_page_break()
        elif btype == "image":
            fig_n += 1
            local, _ = resolve_image_file(block, cwd=cwd, doc_id=doc_id)
            if local and local.is_file():
                try:
                    doc.add_picture(str(local), width=Inches(5.8))
                    last = doc.paragraphs[-1]
                    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    p = doc.add_paragraph()
                    run = p.add_run(f"[图片无法嵌入：{local.name}]")
                    _set_run_font(run, theme, size_pt=10, italic=True, color=theme["muted"])
            else:
                p = doc.add_paragraph()
                run = p.add_run(f"[图片：{block.get('alt') or block.get('path') or block.get('url') or 'missing'}]")
                _set_run_font(run, theme, size_pt=10, italic=True, color=theme["muted"])
            cap = numbered_caption(
                "图",
                fig_n,
                str(block.get("caption") or ""),
                enabled=bool(pack.get("caption_numbers")),
            )
            if cap:
                c = doc.add_paragraph()
                cr = c.add_run(cap)
                _set_run_font(cr, theme, size_pt=9.5, italic=True, color=theme["muted"])
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _p_spacing(c, before=20, after=100, line=1.2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _ppt_rgb(hex_color: str):
    from pptx.util import Pt  # noqa: F401
    from pptx.dml.color import RGBColor

    return RGBColor(*_hex_rgb(hex_color))


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        name = (getattr(layout, "name", "") or "").lower()
        if "blank" in name or "空白" in name:
            return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def _set_slide_bg(slide, hex_color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _ppt_rgb(hex_color)


def _add_rect(slide, left, top, width, height, fill_hex: str):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _ppt_rgb(fill_hex)
    shape.line.fill.background()
    return shape


def _set_ppt_run_font(run, theme: Dict[str, str], *, size: float, bold: bool = False, italic: bool = False, color: str = "", math: bool = False):
    from pptx.util import Pt

    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic or math
    run.font.color.rgb = _ppt_rgb(color or theme["ink"])
    if math:
        run.font.name = "Cambria Math"
    else:
        run.font.name = theme.get("font_heading") if bold else (theme.get("font_body") or "Calibri")


def _clear_ppt_runs(paragraph) -> None:
    p_el = paragraph._p
    for child in list(p_el):
        local = child.tag.split("}")[-1]
        if local in {"r", "br", "fld", "AlternateContent"}:
            p_el.remove(child)


def _fill_ppt_paragraph(
    paragraph,
    text: str,
    theme: Dict[str, str],
    *,
    size: float,
    bold: bool = False,
    italic: bool = False,
    color: str = "",
    prefix: str = "",
):
    # PPT text frames only accept a:r / a:br / a:fld. Injecting mc:AlternateContent
    # + a14:m (and xmlns via lxml .set) produces XML PowerPoint refuses to open.
    _clear_ppt_runs(paragraph)
    if prefix:
        run = paragraph.add_run()
        run.text = prefix
        _set_ppt_run_font(run, theme, size=size, bold=bold, italic=italic, color=color)

    raw = str(text or "")
    if not text_has_math(raw):
        run = paragraph.add_run()
        run.text = raw
        _set_ppt_run_font(run, theme, size=size, bold=bold, italic=italic, color=color)
        return paragraph

    for kind, payload in iter_text_math_parts(raw):
        if kind == "text":
            if not payload:
                continue
            run = paragraph.add_run()
            run.text = payload
            _set_ppt_run_font(run, theme, size=size, bold=bold, italic=italic, color=color)
            continue
        run = paragraph.add_run()
        run.text = latex_to_unicode(payload) or payload
        _set_ppt_run_font(
            run,
            theme,
            size=size + 1,
            bold=False,
            italic=True,
            color=color or theme.get("ink") or "",
            math=True,
        )
    return paragraph


def _textbox(slide, left, top, width, height, text: str, *, theme, size, bold=False, color="", italic=False, align=None):
    from pptx.enum.text import PP_ALIGN

    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align or PP_ALIGN.LEFT
    _fill_ppt_paragraph(p, text, theme, size=size, bold=bold, italic=italic, color=color)
    return box


def _add_notes(slide, notes: str):
    if not notes:
        return
    try:
        ns = slide.notes_slide
        tf = ns.notes_text_frame
        tf.text = notes
    except Exception:
        pass


def _ppt_paper_chrome(slide, theme: Dict[str, str], pack: Dict[str, Any], *, title_bar: bool = False):
    from pptx.util import Inches

    _set_slide_bg(slide, theme["paper"])
    if pack.get("ppt_top_bar"):
        height = 0.14 if title_bar else 0.12
        _add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(height), theme["accent"])


def _ppt_title_rule(slide, theme: Dict[str, str], pack: Dict[str, Any], *, left=0.7, top=1.12):
    from pptx.util import Inches

    if not pack.get("ppt_rule"):
        return
    if pack.get("ppt_top_bar"):
        _add_rect(slide, Inches(left), Inches(top), Inches(1.4), Inches(0.06), theme["accent"])
    else:
        _add_rect(slide, Inches(left), Inches(top), Inches(2.2), Inches(0.02), theme["rule"])


def render_pptx(outline: Dict[str, Any], *, cwd: str = "") -> bytes:
    try:
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise OfficeOutlineError("python-pptx is not installed") from exc

    _sid, theme, pack = _pack_of(outline)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = _blank_layout(prs)
    doc_id = str(outline.get("doc_id") or "")
    slides = list(outline.get("slides") or [])
    total = max(1, len(slides))
    heading_size = float(pack.get("ppt_heading_size") or 26)
    body_size = float(pack.get("ppt_body_size") or 20)
    section_n = 0

    for idx, spec in enumerate(slides, 1):
        if not isinstance(spec, dict):
            continue
        slide = prs.slides.add_slide(layout)
        stype = spec.get("type")
        if stype == "section":
            section_n += 1
            if pack.get("ppt_section_fill") == "dark":
                _set_slide_bg(slide, theme["accent_dark"])
                _add_rect(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), theme["accent"])
                kicker = str(spec.get("kicker") or "SECTION")
                _textbox(
                    slide,
                    Inches(0.9),
                    Inches(2.15),
                    Inches(11.4),
                    Inches(0.45),
                    kicker.upper(),
                    theme=theme,
                    size=13,
                    bold=True,
                    color=theme["header_fg"],
                )
                _textbox(
                    slide,
                    Inches(0.9),
                    Inches(2.7),
                    Inches(11.4),
                    Inches(2.2),
                    str(spec.get("title") or ""),
                    theme=theme,
                    size=36,
                    bold=True,
                    color=theme["header_fg"],
                )
            else:
                _set_slide_bg(slide, theme["paper"])
                _textbox(
                    slide,
                    Inches(0.9),
                    Inches(1.85),
                    Inches(11.4),
                    Inches(0.9),
                    f"{section_n:02d}",
                    theme=theme,
                    size=44,
                    bold=True,
                    color=theme["accent"],
                )
                _textbox(
                    slide,
                    Inches(0.9),
                    Inches(2.85),
                    Inches(11.4),
                    Inches(1.8),
                    str(spec.get("title") or ""),
                    theme=theme,
                    size=32,
                    bold=True,
                    color=theme["ink"],
                )
                _ppt_title_rule(slide, theme, pack, left=0.9, top=4.75)
        elif stype == "title":
            _ppt_paper_chrome(slide, theme, pack, title_bar=True)
            if pack.get("ppt_top_bar"):
                _add_rect(slide, Inches(0.9), Inches(3.05), Inches(2.1), Inches(0.08), theme["accent"])
            else:
                _ppt_title_rule(slide, theme, pack, left=0.9, top=3.15)
            _textbox(
                slide,
                Inches(0.9),
                Inches(2.15),
                Inches(11.4),
                Inches(0.9),
                str(spec.get("title") or outline.get("title") or ""),
                theme=theme,
                size=float(pack.get("ppt_title_size") or 40),
                bold=True,
                color=theme["ink"],
            )
            sub = str(spec.get("subtitle") or outline.get("subtitle") or "")
            if sub:
                _textbox(
                    slide,
                    Inches(0.9),
                    Inches(3.35),
                    Inches(11.4),
                    Inches(1.2),
                    sub,
                    theme=theme,
                    size=18,
                    color=theme["muted"],
                )
        elif stype == "quote":
            _ppt_paper_chrome(slide, theme, pack)
            if pack.get("ppt_quote_bar"):
                _add_rect(slide, Inches(0), Inches(0), Inches(0.16), Inches(7.5), theme["accent"])
            _textbox(
                slide,
                Inches(1.2),
                Inches(2.2),
                Inches(10.8),
                Inches(2.6),
                str(spec.get("text") or ""),
                theme=theme,
                size=26,
                italic=True,
                color=theme["ink"],
            )
            attr = str(spec.get("attribution") or "").strip()
            if attr:
                _textbox(
                    slide,
                    Inches(1.2),
                    Inches(5.0),
                    Inches(10.8),
                    Inches(0.5),
                    f"— {attr}",
                    theme=theme,
                    size=14,
                    color=theme["muted"],
                )
        elif stype == "two_column":
            _ppt_paper_chrome(slide, theme, pack)
            _textbox(
                slide,
                Inches(0.7),
                Inches(0.38),
                Inches(12),
                Inches(0.7),
                str(spec.get("title") or ""),
                theme=theme,
                size=heading_size,
                bold=True,
                color=theme["ink"],
            )
            _ppt_title_rule(slide, theme, pack)
            col_head = theme["accent_dark"] if pack.get("ppt_top_bar") else theme["ink"]
            for col, left in ((spec.get("left") or {}, 0.7), (spec.get("right") or {}, 7.05)):
                heading = str(col.get("heading") or "")
                top = 1.45
                if heading:
                    _textbox(
                        slide,
                        Inches(left),
                        Inches(top),
                        Inches(5.5),
                        Inches(0.45),
                        heading,
                        theme=theme,
                        size=16,
                        bold=True,
                        color=col_head,
                    )
                    top = 1.95
                body = str(col.get("body") or "")
                items = col.get("items") or []
                text = body
                if items:
                    text = (text + "\n" if text else "") + "\n".join(f"•  {it}" for it in items)
                _textbox(
                    slide,
                    Inches(left),
                    Inches(top),
                    Inches(5.5),
                    Inches(4.6),
                    text,
                    theme=theme,
                    size=16,
                    color=theme["ink"],
                )
        elif stype == "image":
            _ppt_paper_chrome(slide, theme, pack)
            title = str(spec.get("title") or "")
            if title:
                _textbox(
                    slide,
                    Inches(0.7),
                    Inches(0.32),
                    Inches(12),
                    Inches(0.55),
                    title,
                    theme=theme,
                    size=22,
                    bold=True,
                    color=theme["ink"],
                )
            local, _ = resolve_image_file(spec, cwd=cwd, doc_id=doc_id)
            if local and local.is_file():
                try:
                    slide.shapes.add_picture(str(local), Inches(2.4), Inches(1.15), width=Inches(8.5))
                except Exception:
                    _textbox(
                        slide,
                        Inches(2.4),
                        Inches(3.0),
                        Inches(8.5),
                        Inches(1),
                        f"[无法嵌入图片 {local.name}]",
                        theme=theme,
                        size=14,
                        italic=True,
                        color=theme["muted"],
                    )
            caption = str(spec.get("caption") or "").strip()
            if caption:
                _textbox(
                    slide,
                    Inches(0.7),
                    Inches(6.55),
                    Inches(12),
                    Inches(0.4),
                    caption,
                    theme=theme,
                    size=13,
                    italic=True,
                    color=theme["muted"],
                    align=PP_ALIGN.CENTER,
                )
        elif stype == "equation":
            _ppt_paper_chrome(slide, theme, pack)
            title = str(spec.get("title") or "")
            top = 0.38
            if title:
                _textbox(
                    slide,
                    Inches(0.7),
                    Inches(top),
                    Inches(12),
                    Inches(0.7),
                    title,
                    theme=theme,
                    size=heading_size,
                    bold=True,
                    color=theme["ink"],
                )
                _ppt_title_rule(slide, theme, pack)
                top = 1.45
            latex = latex_from_block(spec)
            wrapped = latex if text_has_math(latex) else f"${latex}$"
            eq_size = 28 if pack.get("ppt_top_bar") else 24
            _textbox(
                slide,
                Inches(0.9),
                Inches(top + 0.35),
                Inches(11.5),
                Inches(3.6),
                wrapped,
                theme=theme,
                size=eq_size,
                italic=True,
                color=theme["ink"],
                align=PP_ALIGN.CENTER,
            )
            caption = str(spec.get("caption") or "").strip()
            if caption:
                _textbox(
                    slide,
                    Inches(0.7),
                    Inches(6.45),
                    Inches(12),
                    Inches(0.4),
                    caption,
                    theme=theme,
                    size=13,
                    italic=True,
                    color=theme["muted"],
                    align=PP_ALIGN.CENTER,
                )
        else:  # bullets (default)
            _ppt_paper_chrome(slide, theme, pack)
            _textbox(
                slide,
                Inches(0.7),
                Inches(0.38),
                Inches(12),
                Inches(0.7),
                str(spec.get("title") or ""),
                theme=theme,
                size=heading_size,
                bold=True,
                color=theme["ink"],
            )
            _ppt_title_rule(slide, theme, pack)
            items = spec.get("items") or []
            box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(11.8), Inches(5.2))
            tf = box.text_frame
            tf.word_wrap = True
            for i, item in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.level = 0
                p.space_after = Pt(12)
                _fill_ppt_paragraph(p, str(item), theme, size=body_size, color=theme["ink"], prefix="•   ")

        if stype != "section" or pack.get("ppt_footer_on_section"):
            _textbox(
                slide,
                Inches(0.7),
                Inches(7.08),
                Inches(10),
                Inches(0.28),
                str(outline.get("title") or ""),
                theme=theme,
                size=10,
                color=theme["muted"],
            )
            _textbox(
                slide,
                Inches(11.4),
                Inches(7.08),
                Inches(1.3),
                Inches(0.28),
                f"{idx} / {total}",
                theme=theme,
                size=10,
                color=theme["muted"],
                align=PP_ALIGN.RIGHT,
            )
        _add_notes(slide, str(spec.get("notes") or ""))

    buf = io.BytesIO()
    prs.save(buf)
    data = buf.getvalue()
    _ensure_pptx_slides_parse(data)
    return data


def _ensure_pptx_slides_parse(data: bytes) -> None:
    """Refuse to ship slide XML that PowerPoint would repair (xmlns:ns0 / a14:m)."""
    from zipfile import ZipFile

    from lxml import etree

    with ZipFile(io.BytesIO(data)) as zf:
        names = [n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        if not names:
            raise OfficeOutlineError("generated pptx has no slides")
        for name in names:
            raw = zf.read(name)
            if b"xmlns:ns0" in raw or b"ns0:a14" in raw or b"a14:m" in raw:
                raise OfficeOutlineError(
                    f"{name} contains invalid namespace markup (xmlns:ns0 / a14:m)"
                )
            try:
                etree.fromstring(raw)
            except etree.XMLSyntaxError as exc:
                raise OfficeOutlineError(f"{name} is not well-formed XML: {exc}") from exc


def render_outline(outline: Dict[str, Any], *, cwd: str = "") -> bytes:
    kind = outline.get("kind") or "docx"
    if kind == "pptx":
        return render_pptx(outline, cwd=cwd)
    return render_docx(outline, cwd=cwd)
